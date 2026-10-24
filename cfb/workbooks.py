"""Phase 4: Excel workbooks (schedule + postseason). Round labels abbreviated per spec."""
import pickle, collections
import sys; sys.path.insert(0,'/home/claude/cfb2')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pbp_engine as E
U = E.U

F="Arial"; HF=PatternFill("solid",fgColor="14213D"); HFO=Font(name=F,bold=True,color="FFFFFF",size=11)
TT=Font(name=F,bold=True,size=14); N=Font(name=F,size=10)
thin=Side(style="thin",color="D9D9D9"); B=Border(left=thin,right=thin,top=thin,bottom=thin)

def schedule_workbook(out='/mnt/user-data/outputs/Season1_Schedule.xlsx'):
    with open('/home/claude/cfb2/state/schedule_s1.pkl','rb') as f:
        sched = pickle.load(f)
    wb = Workbook(); ws = wb.active; ws.title = "Full Schedule"
    heads = ["Wk","Home","Away","Type","Locked?"]
    ws.append(heads)
    for i,h in enumerate(heads,1):
        c=ws.cell(row=1,column=i); c.font=HFO; c.fill=HF; c.alignment=Alignment(horizontal="center")
    locked = {frozenset((p["team_a"],p["team_b"])) for p in U["fixed_pairs"]}
    for wk,h,a,ic in sched:
        tag = "RIVALRY" if (frozenset((h,a)) in locked and wk==14) else ("FIXED" if frozenset((h,a)) in locked else "")
        ws.append([wk,h,a,"Conf" if ic else "NC",tag])
    for r in range(2,ws.max_row+1):
        for c in range(1,6):
            cell=ws.cell(row=r,column=c); cell.font=N; cell.border=B
    for i,w in enumerate([6,22,22,8,10],1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:E{ws.max_row}"
    wb.save(out); return out

def postseason_workbook(rounds, out='/mnt/user-data/outputs/Season1_Postseason.xlsx'):
    """rounds: list of (abbrev_label, [ (seedH, teamH, scoreH, seedA, teamA, scoreA, venue) ])"""
    wb = Workbook(); ws = wb.active; ws.title = "Postseason"
    r = 1
    for label, games in rounds:
        ws.cell(row=r,column=1,value=label).font = TT; r += 1
        heads = ["Seed","Team","Pts","Seed","Team","Pts","Venue"]
        for i,h in enumerate(heads,1):
            c=ws.cell(row=r,column=i,value=h); c.font=HFO; c.fill=HF
        r += 1
        for g in games:
            for c,v in enumerate(g,1):
                cell=ws.cell(row=r,column=c,value=v); cell.font=N; cell.border=B
            r += 1
        r += 1
    for i,w in enumerate([6,22,6,6,22,6,26],1): ws.column_dimensions[get_column_letter(i)].width=w
    wb.save(out); return out
